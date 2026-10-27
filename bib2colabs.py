#!/usr/bin/env python3

r"""
bib2colabs.py

Merge BibTeX, COA workbook, and saved database data.

Typical workflow for updating the COA:

    1. Update the .bib files in data/ (graduate-committee entries must have
       keywords = {graduate-committee}).
    2. Run, using your latest COA workbook as the base:

        python bib2colabs.py data --coa 20261027-Colbry_COA.xlsx \
            --database-out coa_database.csv --coa-out coa.csv \
            --table3-out table3.csv --years 4

    3. Copy/paste into the new COA, cleaning up as needed:
        coa.csv     -> Table 4 (A:/C:/T: rows active within --years; delete
                       the T: rows, which belong in Table 3)
        table3.csv  -> Table 3 (existing G:/T: rows plus new T: people)
       Check for the same person spelled differently (e.g. names that appear
       in both the COA and the .bib files) since names match exactly.
    4. coa_database.csv keeps everything (all years, bibkeys, flags). Reuse
       it later with --database-in instead of re-reading everything:

        python bib2colabs.py --database-in coa_database.csv \
            --coa-out coa.csv --years 4

Other notes:

    Inputs (any combination): a .bib file or folder, --coa workbook, and
    --database-in. Outputs: --coa-out, --table3-out, --database-out.
    --years filters only coa.csv; "Present" and unknown years are kept.
    Use --me NAME ... to exclude your own author names.

Requirements:

    pip install pandas bibtexparser openpyxl
"""

from collections import defaultdict
from datetime import datetime
import argparse
import re
from pathlib import Path

import pandas as pd
import bibtexparser
from openpyxl import load_workbook


def normalize_author(author):
    """
    Normalize author names for duplicate detection.

    This is intentionally conservative. It will merge:

        "Smith, John"
        "smith, john"
        " Smith,   John "

    but not:

        "Smith, John"
        "Smith, J."
    """

    author = author.strip().lower()
    author = re.sub(r"\s+", " ", author)

    return author


def bib_to_records(bibfile):
    """
    Convert a BibTeX file into a publication-author dataframe.

    Returns one row per publication-author pair.

    Columns:
        author
        year
        title
        bibkey
        committee (True for entries keyworded ``graduate-committee``)
    """

    with open(bibfile, "r", encoding="utf-8") as bibtex_file:
        bib_database = bibtexparser.load(bibtex_file)

    rows = []

    for entry in bib_database.entries:

        bibkey = entry.get("ID", "<UNKNOWN KEY>")

        try:

            if "author" not in entry:
                raise ValueError("missing author field")

            year_value = entry.get("year")
            if year_value is None:
                date_match = re.match(
                    r"\s*(\d{4})",
                    str(entry.get("date", "")),
                )
                if date_match is None:
                    raise ValueError("missing year field or usable date field")
                year_value = date_match.group(1)

            try:
                year = int(year_value)
            except Exception:
                year_match = re.search(r"\b\d{4}\b", str(year_value))
                if year_match is None:
                    raise ValueError(f"invalid year value: {year_value}")
                year = int(year_match.group())

            authors = [
                author.strip()
                for author in entry["author"].split(" and ")
            ]

            title = entry.get("title", "")
            committee = bool(
                re.search(
                    r"graduate[-_ ]committee",
                    str(entry.get("keywords", "")),
                    re.IGNORECASE,
                )
            )

            for author in authors:
                rows.append(
                    {
                        "author": author,
                        "year": year,
                        "title": title,
                        "bibkey": bibkey,
                        "committee": committee,
                    }
                )

        except Exception as e:

            print(
                f"ERROR processing BibTeX entry '{bibkey}': {e}"
            )

            print("Entry contents:")
            print(entry)
            print("-" * 80)

            raise

    return pd.DataFrame(
        rows,
        columns=["author", "year", "title", "bibkey", "committee"],
    )


def collaborators_from_records(records_df, my_names=None):
    """
    Generate collaborator statistics from publication records.

    Parameters
    ----------
    records_df : pandas.DataFrame
    my_names : list[str]
        Author names to exclude.

    Returns
    -------
    pandas.DataFrame
    """

    if my_names is None:
        my_names = []

    collaborators = defaultdict(
        lambda: {
            "n_publications": 0,
            "first_year": None,
            "last_year": None,
            "bibkeys": [],
            "graduate_committee": False,
        }
    )

    for _, row in records_df.iterrows():

        author = row["author"]

        if author in my_names:
            continue

        year = row["year"]

        collaborators[author]["n_publications"] += 1
        if "committee" in row and row["committee"] is True:
            collaborators[author]["graduate_committee"] = True
        if "bibkey" in row and pd.notna(row["bibkey"]):
            collaborators[author]["bibkeys"].append(str(row["bibkey"]))

        if collaborators[author]["first_year"] is None:
            collaborators[author]["first_year"] = year
        else:
            collaborators[author]["first_year"] = min(
                collaborators[author]["first_year"],
                year,
            )

        if collaborators[author]["last_year"] is None:
            collaborators[author]["last_year"] = year
        else:
            collaborators[author]["last_year"] = max(
                collaborators[author]["last_year"],
                year,
            )

    rows = []

    for author, data in collaborators.items():

        rows.append(
            {
                "author": author,
                "n_publications": data["n_publications"],
                "first_year": data["first_year"],
                "last_year": data["last_year"],
                "bibkeys": ", ".join(sorted(set(data["bibkeys"]))),
                "graduate_committee": data["graduate_committee"],
            }
        )

    if not rows:
        return pd.DataFrame()

    return pd.DataFrame(rows)


def merge_duplicate_collaborators(df):
    """
    Merge collaborators with identical normalized names.

    Uses:
        normalize_author()

    Returns a cleaned dataframe.
    """

    if df.empty:
        return df

    df = df.copy()

    df["author_normalized"] = (
        df["author"]
        .astype(str)
        .apply(normalize_author)
    )

    merged = (
        df.groupby("author_normalized", as_index=False)
        .agg(
            author=("author", "first"),
            n_publications=("n_publications", "sum"),
            first_year=("first_year", "min"),
            last_year=("last_year", "max"),
            graduate_committee=("graduate_committee", "any"),
            bibkeys=(
                "bibkeys",
                lambda values: ", ".join(
                    sorted(
                        {
                            key.strip()
                            for value in values
                            if pd.notna(value)
                            for key in str(value).split(",")
                            if key.strip()
                        }
                    )
                ),
            ),
        )
    )

    merged = merged.sort_values(
        ["n_publications", "author"],
        ascending=[False, True],
    )

    return merged.drop(columns=["author_normalized"])


def bib_to_collaborators(
    bibfile,
    my_names=None,
):
    """
    Convenience wrapper.

    Reads BibTeX and returns cleaned collaborator dataframe.
    """

    records = bib_to_records(bibfile)

    collaborators = collaborators_from_records(
        records,
        my_names=my_names,
    )

    collaborators = merge_duplicate_collaborators(
        collaborators
    )

    return collaborators


def read_coa_tables(coa_file):
    """Read the five COA workbook tables into a dictionary of DataFrames.

    The result uses ``table1`` through ``table5`` as keys. Each DataFrame has
    descriptive column names, and tables with category markers retain those
    markers in an ``indicator`` column. Empty tables are returned as empty
    DataFrames with their expected columns.
    """
    table_schemas = {
        1: (("name", 1), ("affiliation", 2), ("last_active", 3)),
        2: (
            ("indicator", 0),
            ("name", 1),
            ("relationship", 2),
            ("email", 3),
            ("last_active", 4),
        ),
        3: (
            ("indicator", 0),
            ("name", 1),
            ("affiliation", 2),
            ("email", 3),
        ),
        4: (
            ("indicator", 0),
            ("name", 1),
            ("affiliation", 2),
            ("email", 3),
            ("last_active", 4),
        ),
        5: (
            ("indicator", 0),
            ("name", 1),
            ("affiliation", 2),
            ("journal_collection", 3),
            ("last_active", 4),
        ),
    }
    valid_indicators = {
        2: {"R:"},
        3: {"G:", "T:"},
        4: {"A:", "C:"},
        5: {"B:", "E:"},
    }

    workbook = load_workbook(coa_file, data_only=True, read_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    table_headers = {}
    for row_index, row in enumerate(rows):
        if not row or row[0] is None:
            continue

        header_value = str(row[0]).strip()
        if header_value in {"1", "2", "3", "4", "5"}:
            table_number = int(header_value)
            table_headers[table_number] = row_index

    tables = {}
    for table_number, schema in table_schemas.items():
        if table_number not in table_headers:
            raise ValueError(
                f"Could not find the header row for COA Table {table_number}"
            )

        header_index = table_headers[table_number]
        section_end = next(
            (
                row_index
                for row_index in range(header_index + 1, len(rows))
                if rows[row_index]
                and rows[row_index][0] is not None
                and str(rows[row_index][0]).strip().startswith("Table ")
            ),
            len(rows),
        )

        records = []
        for row in rows[header_index + 1:section_end]:
            record = {
                column: (
                    row[column_index].strip()
                    if isinstance(row[column_index], str)
                    else row[column_index]
                )
                for column, column_index in schema
                if column_index < len(row)
            }

            name = record.get("name")
            if name is None or not str(name).strip():
                continue

            indicator = record.get("indicator")
            if (
                table_number in valid_indicators
                and indicator not in valid_indicators[table_number]
            ):
                continue

            records.append(record)

        tables[f"table{table_number}"] = pd.DataFrame(
            records,
            columns=[column for column, _ in schema],
        )

    return tables


def merge_coa_sources(coa_tables, bib_collaborators, database_df=None):
    """Merge COA Table 4 records with collaborators summarized from BibTeX.

    Names are matched case- and whitespace-insensitively. For matched people,
    keep populated COA names, indicators, affiliations, and email/contact
    details, while setting ``last_active`` to the more recent COA/BibTeX year.
    BibTeX-only people receive the ``A:`` indicator, except people on
    ``graduate-committee`` entries, who are flagged ``graduate_committee`` and
    get ``T:``, overriding ``A:``/``C:`` (for COA Table 3). Publication counts,
    first/last publication years, and comma-separated BibTeX keys are retained
    as additional columns.

    Existing database rows are retained and enriched by new COA/BibTeX data.
    For authors already in the database, publication counts are not added a
    second time; the larger count and union of known BibTeX keys are retained.

    Returns a new DataFrame; input DataFrames are not modified.
    """
    standard_columns = [
        "indicator",
        "name",
        "affiliation",
        "email",
        "last_active",
        "n_publications",
        "first_year",
        "last_year",
        "bibkeys",
        "graduate_committee",
    ]
    output_columns = standard_columns + (
        [
            column
            for column in database_df.columns
            if column not in standard_columns
        ]
        if database_df is not None
        else []
    )
    coa_columns = ["indicator", "name", "affiliation", "email", "last_active"]

    if database_df is not None:
        missing_columns = [
            column
            for column in ("name", "indicator")
            if column not in database_df.columns
        ]
        if missing_columns:
            raise ValueError(
                "Database is missing required columns: "
                + ", ".join(missing_columns)
            )

    if coa_tables is None:
        coa_df = pd.DataFrame(columns=coa_columns)
    else:
        if "table4" not in coa_tables:
            raise ValueError("COA tables must include a 'table4' DataFrame")
        coa_df = coa_tables["table4"]
        if "name" not in coa_df.columns:
            raise ValueError("COA Table 4 must include a 'name' column")

    if not bib_collaborators.empty and "author" not in bib_collaborators.columns:
        raise ValueError("BibTeX collaborators must include an 'author' column")

    def is_missing(value):
        return value is None or (
            isinstance(value, str) and not value.strip()
        ) or bool(pd.isna(value))

    def recency(value):
        if is_missing(value):
            return (0, 0)
        if isinstance(value, str) and value.strip().lower() == "present":
            return (2, 0)
        if hasattr(value, "year"):
            return (1, int(value.year))
        try:
            return (1, int(value))
        except (TypeError, ValueError):
            return (0, 0)

    def latest_value(current, candidate):
        return candidate if recency(candidate) > recency(current) else current

    def combine_bibkeys(values):
        bibkeys = set()
        for value in values:
            if not is_missing(value):
                bibkeys.update(
                    bibkey.strip()
                    for bibkey in str(value).split(",")
                    if bibkey.strip()
                )
        return ", ".join(sorted(bibkeys))

    def empty_record():
        record = {
            "indicator": None,
            "name": None,
            "affiliation": None,
            "email": None,
            "last_active": None,
            "n_publications": 0,
            "first_year": None,
            "last_year": None,
            "bibkeys": "",
            "graduate_committee": False,
        }
        record.update(
            {column: None for column in output_columns if column not in record}
        )
        return record

    merged = {}
    if database_df is not None:
        for _, row in database_df.iterrows():
            name = row.get("name")
            if is_missing(name):
                continue

            key = normalize_author(str(name))
            record = merged.setdefault(key, empty_record())
            for column in output_columns:
                value = row.get(column)
                if column == "last_active":
                    record[column] = latest_value(record[column], value)
                elif column == "n_publications":
                    if not is_missing(value):
                        record[column] = max(record[column], int(value))
                elif column == "first_year":
                    if not is_missing(value):
                        record[column] = (
                            int(value)
                            if is_missing(record[column])
                            else min(record[column], int(value))
                        )
                elif column == "last_year":
                    if not is_missing(value):
                        record[column] = (
                            int(value)
                            if is_missing(record[column])
                            else max(record[column], int(value))
                        )
                elif column == "bibkeys":
                    record[column] = combine_bibkeys(
                        [record[column], value]
                    )
                elif column == "graduate_committee":
                    record[column] = bool(record[column]) or (
                        not is_missing(value)
                        and str(value).strip().lower() in ("true", "1")
                    )
                elif is_missing(record[column]) and not is_missing(value):
                    record[column] = value
                elif column not in standard_columns:
                    record[column] = value

    for _, row in coa_df.iterrows():
        name = row.get("name")
        if is_missing(name):
            continue

        key = normalize_author(str(name))
        record = merged.setdefault(key, empty_record())
        for column in coa_columns:
            value = row.get(column)
            if column == "last_active":
                record[column] = latest_value(record[column], value)
            elif column == "indicator" and record[column] == "T:":
                continue
            elif not is_missing(value):
                record[column] = value

    bib_df = bib_collaborators.copy()
    if bib_df.empty:
        bib_groups = []
    else:
        bib_df["_author_key"] = bib_df["author"].map(
            lambda name: normalize_author(str(name))
        )
        bib_groups = bib_df.groupby("_author_key", sort=False)

    for key, group in bib_groups:
        author = next(
            (name for name in group["author"] if not is_missing(name)),
            None,
        )
        if author is None:
            continue

        record = merged.setdefault(key, empty_record())

        if is_missing(record["indicator"]):
            record["indicator"] = "A:"
        if (
            "graduate_committee" in group.columns
            and group["graduate_committee"].any()
        ):
            record["graduate_committee"] = True
        if is_missing(record["name"]):
            record["name"] = author

        if "n_publications" in group.columns:
            publication_count = int(
                pd.to_numeric(group["n_publications"], errors="coerce")
                .fillna(0)
                .sum()
            )
            if database_df is not None:
                record["n_publications"] = max(
                    record["n_publications"],
                    publication_count,
                )
            else:
                record["n_publications"] += publication_count
        for year_column, operation in (
            ("first_year", min),
            ("last_year", max),
        ):
            if year_column in group.columns:
                years = pd.to_numeric(group[year_column], errors="coerce").dropna()
                if not years.empty:
                    value = operation(years)
                    if is_missing(record[year_column]):
                        record[year_column] = int(value)
                    elif year_column == "first_year":
                        record[year_column] = min(record[year_column], int(value))
                    else:
                        record[year_column] = max(record[year_column], int(value))

        if "last_year" in group.columns:
            years = pd.to_numeric(group["last_year"], errors="coerce").dropna()
            if not years.empty:
                record["last_active"] = latest_value(
                    record["last_active"], int(years.max())
                )
        if "bibkeys" in group.columns:
            record["bibkeys"] = combine_bibkeys(
                [record["bibkeys"], *group["bibkeys"].tolist()]
            )

    for record in merged.values():
        if record["graduate_committee"]:
            record["indicator"] = "T:"

    result = pd.DataFrame(merged.values(), columns=output_columns)
    if not result.empty:
        result = result.sort_values("name", kind="stable").reset_index(drop=True)
    return result


def coa_out(merged_df, csv_file, years=3):
    """Write merged collaborator data as a COA-ready CSV.

    Select and format only the COA Table 4 fields for copy/paste: the
    author/collaborator indicator, name, organizational affiliation,
    email/department, and last active year. Keep supplemental fields such as
    publication counts, year range, and titles in ``merged_df`` for other uses;
    do not mutate the input dataframe when preparing the COA-only CSV. Apply
    the recency filter here, so database exports retain all years.
    """
    coa_columns = [
        "indicator",
        "name",
        "affiliation",
        "email",
        "last_active",
    ]
    missing_columns = [column for column in coa_columns if column not in merged_df]
    if missing_columns:
        raise ValueError(
            "Merged dataframe is missing COA columns: "
            + ", ".join(missing_columns)
        )

    if years < 0:
        raise ValueError("years must be zero or greater")

    min_year = datetime.now().year - years
    last_active_year = pd.to_numeric(
        merged_df["last_active"], errors="coerce"
    )
    is_present = (
        merged_df["last_active"]
        .astype("string")
        .str.strip()
        .str.casefold()
        .eq("present")
    )
    is_recent = last_active_year.ge(min_year) | is_present | last_active_year.isna()

    merged_df.loc[is_recent, coa_columns].to_csv(csv_file, index=False)


def table3_out(merged_df, coa_tables, csv_file):
    """Write COA Table 3 (advisors/advisees) as a CSV.

    Combine existing Table 3 rows from ``coa_tables`` (may be None) with the
    ``T:`` rows in ``merged_df``. People are matched by normalized name; COA
    values are kept and only missing affiliation/email are filled from the
    merged data. No year filter is applied. Returns the Table 3 DataFrame.
    """
    columns = ["indicator", "name", "affiliation", "email"]

    def clean(value):
        return None if pd.isna(value) or not str(value).strip() else value

    rows = {}
    if coa_tables is not None and "table3" in coa_tables:
        for _, row in coa_tables["table3"].iterrows():
            if clean(row.get("name")) is None:
                continue
            rows.setdefault(
                normalize_author(str(row["name"])),
                {column: clean(row.get(column)) for column in columns},
            )

    committee = merged_df[merged_df["indicator"] == "T:"]
    for _, row in committee.iterrows():
        record = rows.setdefault(
            normalize_author(str(row["name"])),
            {column: None for column in columns},
        )
        record["indicator"] = record["indicator"] or "T:"
        for column in ("name", "affiliation", "email"):
            if record[column] is None:
                record[column] = clean(row.get(column))

    table3 = pd.DataFrame(rows.values(), columns=columns)
    table3.to_csv(csv_file, index=False)
    return table3


def to_database(merged_df, database_file):
    """Write all merged COA and BibTeX columns to a database CSV."""
    merged_df.to_csv(database_file, index=False)


def from_database(database_file):
    """Load a database CSV, retaining all stored columns.

    The CSV must contain at least ``indicator`` and ``name`` columns, as
    written by ``to_database``. Legacy publication titles are retained under
    ``legacy_titles`` and a legacy ``optional`` column is renamed to ``email``
    when loading databases created by earlier versions.
    """
    database_df = pd.read_csv(database_file)
    missing_columns = [
        column
        for column in ("indicator", "name")
        if column not in database_df.columns
    ]
    if missing_columns:
        raise ValueError(
            "Database CSV is missing required columns: "
            + ", ".join(missing_columns)
        )
    if "titles" in database_df.columns:
        database_df = database_df.rename(columns={"titles": "legacy_titles"})
    if "optional" in database_df.columns and "email" not in database_df.columns:
        database_df = database_df.rename(columns={"optional": "email"})
    return database_df


def generate_coa_csv(
    bib_source=None,
    csv_file=None,
    coa_file=None,
    years=3,
    my_names=None,
    database_file=None,
    database_out=None,
    table3_file=None,
):
    """Merge optional BibTeX, COA, and database inputs, then export outputs.

    ``bib_source`` may be one ``.bib`` file or a folder of them. A saved
    database can be loaded and updated without re-reading prior inputs.
    ``table3_file`` writes Table 3: COA Table 3 rows plus new ``T:`` entries.
    ``csv_file`` is the COA-only CSV output; ``database_out`` writes the full
    merged DataFrame. Both outputs may be written in one run. ``years`` filters
    only the COA CSV output; the returned DataFrame and database CSV retain all
    years of available information.
    """
    if csv_file is None and database_out is None and table3_file is None:
        raise ValueError("Specify at least one output")
    if bib_source is None and coa_file is None and database_file is None:
        raise ValueError("Specify at least one input source")

    database_df = (
        from_database(database_file)
        if database_file is not None
        else None
    )

    if bib_source is None:
        collaborators = pd.DataFrame(
            columns=[
                "author",
                "n_publications",
                "first_year",
                "last_year",
                "bibkeys",
            ]
        )
    else:
        source_path = Path(bib_source)
        if source_path.is_dir():
            bib_files = sorted(source_path.glob("*.bib"))
            if not bib_files:
                raise ValueError(f"No .bib files found in folder: {source_path}")
        elif source_path.is_file():
            bib_files = [source_path]
        else:
            raise FileNotFoundError(
                f"BibTeX input does not exist: {source_path}"
            )

        records = pd.concat(
            [bib_to_records(bib_file) for bib_file in bib_files],
            ignore_index=True,
        )
        if not records.empty:
            records["_author_key"] = records["author"].map(
                lambda name: normalize_author(str(name))
            )
            records["_title_key"] = records["title"].map(
                lambda title: re.sub(r"\s+", " ", str(title).strip().lower())
            )
            records = records.drop_duplicates(
                subset=["_author_key", "year", "_title_key"]
            ).drop(columns=["_author_key", "_title_key"])
        collaborators = collaborators_from_records(
            records,
            my_names=my_names,
        )
        collaborators = merge_duplicate_collaborators(collaborators)

    coa_tables = read_coa_tables(coa_file) if coa_file is not None else None
    merged = merge_coa_sources(
        coa_tables,
        collaborators,
        database_df=database_df,
    )

    if csv_file is not None:
        coa_out(merged, csv_file, years=years)
    if table3_file is not None:
        table3_out(merged, coa_tables, table3_file)
    if database_out is not None:
        to_database(merged, database_out)
    return merged


if __name__ == "__main__":

    parser = argparse.ArgumentParser(
        description="Merge BibTeX, COA, and collaborator database data.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )

    parser.add_argument(
        "bib_source",
        nargs="?",
        help="Input BibTeX file or folder containing .bib files",
    )

    parser.add_argument(
        "csvfile",
        nargs="?",
        help="Output COA CSV file",
    )

    parser.add_argument(
        "--years",
        type=int,
        default=3,
        help="COA CSV recency window in years (default=3)",
    )

    parser.add_argument(
        "--me",
        nargs="+",
        default=[],
        help="Author names to exclude",
    )
    parser.add_argument(
        "--coa",
        dest="coa_file",
        help="Optional input COA workbook",
    )
    parser.add_argument(
        "--database-in",
        dest="database_file",
        help="Optional input database CSV to update or export",
    )
    parser.add_argument(
        "--database-out",
        help="Optional output CSV containing all merged information",
    )
    parser.add_argument(
        "--coa-out",
        dest="coa_output",
        help="Output COA-only CSV file (alternative to the positional output)",
    )

    parser.add_argument(
        "--table3-out",
        help="Output CSV for COA Table 3 (COA Table 3 plus new T: entries)",
    )

    args = parser.parse_args()

    try:
        if args.csvfile is not None and args.coa_output is not None:
            parser.error("Use either the positional COA output or --coa-out")
        df = generate_coa_csv(
            args.bib_source,
            args.coa_output or args.csvfile,
            coa_file=args.coa_file,
            years=args.years,
            my_names=args.me,
            database_file=args.database_file,
            database_out=args.database_out,
            table3_file=args.table3_out,
        )
    except (ValueError, FileNotFoundError) as error:
        parser.error(str(error))


    print(
        f"Processed {len(df)} collaborator records"
    )